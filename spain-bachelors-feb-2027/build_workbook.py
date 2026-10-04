"""Build spanish_bachelors_feb_2027.xlsx from the agents' JSON research files.

Reads every *.json in RESEARCH_DIR (one per research group), merges programs,
rejected entries and sources, de-duplicates URLs, sorts programs by the cheapest
deposit needed to obtain an admission letter, and writes three sheets:
Programs, Rejected, Sources.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

RESEARCH_DIR = Path(sys.argv[1])
OUT = Path(sys.argv[2])
ACCESSED = "2026-10-04"

PROGRAM_COLS = [
    ("Rank", 6, None),
    ("University", 28, "university"),
    ("City", 14, "city"),
    ("Public / Private", 22, "type"),
    ("Degree", 34, "degree"),
    ("Field", 12, "field"),
    ("Language", 16, "language"),
    ("Feb 2027 start for NEW students", 14, "feb_start_confirmed"),
    ("Feb-start evidence", 40, "feb_start_evidence"),
    ("Application deadline (Feb 2027 intake)", 22, "application_deadline"),
    ("Application fee (EUR)", 12, "application_fee_eur"),
    ("Application fee detail", 30, "application_fee_text"),
    ("Deposit before admission letter (EUR)", 14, "deposit_before_letter_eur"),
    ("Deposit detail", 40, "deposit_text"),
    ("Deposit deducted from tuition?", 12, "deposit_deducted_from_tuition"),
    ("Refunded if visa refused?", 12, "refund_if_visa_refused"),
    ("Refund conditions", 40, "refund_text"),
    ("Annual tuition non-EU (EUR)", 14, "annual_tuition_non_eu_eur"),
    ("Tuition detail", 40, "tuition_text"),
    ("Foreign bachelor's holders eligible?", 12, "foreign_graduate_eligible"),
    ("Admission route for graduates", 40, "graduate_route"),
    ("Language requirement", 32, "language_requirement"),
    ("Official grado (RUCT)?", 20, "official_ruct"),
    ("Data year", 18, "data_year"),
    ("Page last updated", 14, "page_last_updated"),
    ("Visa-timing note", 36, "visa_timing_note"),
    ("Risks / flags", 40, "risks"),
    ("Sources (by fact)", 60, "_sources"),
    ("Research group", 18, "_group"),
]

REJECTED_COLS = [
    ("University", 30, "university"),
    ("City", 14, "city"),
    ("Public / Private", 22, "type"),
    ("Degree / programme", 34, "degree"),
    ("Reason excluded", 60, "reason"),
    ("Source URL(s)", 70, "_sources"),
    ("Research group", 18, "_group"),
]

SOURCE_COLS = [
    ("#", 5, None),
    ("URL", 70, "url"),
    ("Page title", 40, "title"),
    ("Official source?", 10, "official"),
    ("Page last updated", 16, "last_updated"),
    ("Used for", 50, "used_for"),
    ("Accessed", 12, "_accessed"),
    ("Research group", 18, "_group"),
]

LANG_OK = re.compile(r"english|ingl[eé]s|angl[eè]s|french|franc[eé]s|fran[cç]ais", re.I)

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
OLD_FILL = PatternFill("solid", fgColor="FFF2CC")  # possibly outdated data
RISK_FILL = PatternFill("solid", fgColor="FCE4D6")  # letter not obtainable in time


def as_number(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return None


def deposit_sort_key(p):
    """Cheapest-known first; then unknown deposits; Feb-UNCLEAR rows last."""
    unclear = 0 if str(p.get("feb_start_confirmed", "")).upper().startswith("YES") else 1
    dep = as_number(p.get("deposit_before_letter_eur"))
    fee = as_number(p.get("application_fee_eur")) or 0.0
    tuition = as_number(p.get("annual_tuition_non_eu_eur"))
    if dep is None:
        return (unclear, 1, 0.0, tuition if tuition is not None else 1e9)
    return (unclear, 0, dep + fee, tuition if tuition is not None else 1e9)


def fmt_sources(fs):
    if not fs:
        return ""
    if isinstance(fs, list):
        return "\n".join(fs)
    lines = []
    for k, urls in fs.items():
        if isinstance(urls, str):
            urls = [urls]
        for u in urls or []:
            lines.append(f"[{k}] {u}")
    return "\n".join(lines)


def style_sheet(ws, cols, n_rows, table_name):
    for i, (name, width, _) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "C2" if table_name == "Programs" else "B2"
    if n_rows:
        ref = f"A1:{get_column_letter(len(cols))}{n_rows + 1}"
        t = Table(displayName=table_name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)
    for row in ws.iter_rows(min_row=2, max_row=n_rows + 1):
        for c in row:
            c.alignment = WRAP


def main():
    programs, rejected, sources = [], [], {}
    for f in sorted(RESEARCH_DIR.glob("*.json")):
        data = json.loads(f.read_text(encoding="utf-8"))
        group = data.get("group") or f.stem
        for p in data.get("programs", []):
            p["_group"] = group
            programs.append(p)
        for r in data.get("rejected", []):
            r["_group"] = group
            rejected.append(r)
        for s in data.get("sources", []):
            url = (s.get("url") or "").strip()
            if not url:
                continue
            if url in sources:
                prev = sources[url]
                if s.get("used_for") and s["used_for"] not in (prev.get("used_for") or ""):
                    prev["used_for"] = (prev.get("used_for", "") + "; " + s["used_for"]).strip("; ")
            else:
                s["_group"] = group
                sources[url] = s
        # make sure every URL cited in a row also appears in Sources
        for p in data.get("programs", []):
            fs = p.get("field_sources") or {}
            for k, urls in (fs.items() if isinstance(fs, dict) else []):
                for u in ([urls] if isinstance(urls, str) else urls or []):
                    sources.setdefault(u, {"url": u, "title": "", "official": "",
                                           "last_updated": "", "used_for": f"{p.get('university')}: {k}",
                                           "_group": group})
        for r in data.get("rejected", []):
            for u in r.get("source_urls", []) or []:
                sources.setdefault(u, {"url": u, "title": "", "official": "", "last_updated": "",
                                       "used_for": f"Rejected: {r.get('university')}", "_group": group})

    # The applicant can only study in English or French: anything else is rejected.
    kept = []
    for p in programs:
        if LANG_OK.search(str(p.get("language", ""))):
            kept.append(p)
        else:
            rejected.append({
                "university": p.get("university", ""), "city": p.get("city", ""),
                "type": p.get("type", ""), "degree": p.get("degree", ""),
                "reason": f"Taught in {p.get('language') or 'an unconfirmed language'} — "
                          "applicant needs English or French. (Feb 2027 start: "
                          f"{p.get('feb_start_confirmed', '?')})",
                "source_urls": [u for urls in (p.get("field_sources") or {}).values()
                                for u in ([urls] if isinstance(urls, str) else urls or [])][:3],
                "_group": p["_group"],
            })
    programs = kept
    programs.sort(key=deposit_sort_key)

    wb = Workbook()
    ws = wb.active
    ws.title = "Programs"
    ws.append([c[0] for c in PROGRAM_COLS])
    for rank, p in enumerate(programs, start=1):
        row = []
        for name, _, key in PROGRAM_COLS:
            if key is None:
                row.append(rank)
            elif key == "_sources":
                row.append(fmt_sources(p.get("field_sources")))
            else:
                v = p.get(key, "")
                if isinstance(v, (list, dict)):
                    v = json.dumps(v, ensure_ascii=False)
                row.append(v if v != "" else "NOT FOUND" if key.endswith(("_eur", "deadline")) else v)
        ws.append(row)
        r = ws.max_row
        if re.search(r"2025-26|2024-25|outdated", str(p.get("data_year", "")), re.I):
            for c in ws[r]:
                c.fill = OLD_FILL
        if re.search(r"\b(no|impossible|not realistic|unlikely)\b", str(p.get("visa_timing_note", "")), re.I):
            ws.cell(row=r, column=[c[0] for c in PROGRAM_COLS].index("Visa-timing note") + 1).fill = RISK_FILL
    for col in ("Application fee (EUR)", "Deposit before admission letter (EUR)", "Annual tuition non-EU (EUR)"):
        idx = [c[0] for c in PROGRAM_COLS].index(col) + 1
        for r in range(2, ws.max_row + 1):
            c = ws.cell(row=r, column=idx)
            if isinstance(c.value, (int, float)):
                c.number_format = '#,##0" €"'
    style_sheet(ws, PROGRAM_COLS, len(programs), "Programs")

    ws2 = wb.create_sheet("Rejected")
    ws2.append([c[0] for c in REJECTED_COLS])
    rejected.sort(key=lambda r: (str(r.get("university", "")).lower()))
    for r in rejected:
        ws2.append([fmt_sources(r.get("source_urls")) if k == "_sources" else r.get(k, "")
                    for _, _, k in REJECTED_COLS])
    style_sheet(ws2, REJECTED_COLS, len(rejected), "Rejected")

    ws3 = wb.create_sheet("Sources")
    ws3.append([c[0] for c in SOURCE_COLS])
    for i, s in enumerate(sorted(sources.values(), key=lambda s: s["url"]), start=1):
        row = []
        for _, _, k in SOURCE_COLS:
            if k is None:
                row.append(i)
            elif k == "_accessed":
                row.append(ACCESSED)
            else:
                v = s.get(k, "")
                row.append("Yes" if v is True else "No" if v is False else v)
        ws3.append(row)
        ws3.cell(row=ws3.max_row, column=2).hyperlink = s["url"]
    style_sheet(ws3, SOURCE_COLS, len(sources), "Sources")

    wb.save(OUT)
    print(f"programs={len(programs)} rejected={len(rejected)} sources={len(sources)} -> {OUT}")


if __name__ == "__main__":
    main()
